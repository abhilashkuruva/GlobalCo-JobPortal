import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Palette
COLOR_PRIMARY_HEX = "1E3A8A"      # Deep Navy Blue
COLOR_SECONDARY_HEX = "2563EB"    # Royal Blue
COLOR_ACCENT_HEX = "0D9488"       # Teal
COLOR_DARK_HEX = "0F172A"         # Slate 900
COLOR_MUTED_HEX = "475569"        # Slate 600
COLOR_BORDER_HEX = "CBD5E1"       # Slate 300
COLOR_BG_LIGHT_HEX = "F8FAFC"     # Slate 50
COLOR_CALLOUT_INFO_BG = "EFF6FF"  # Blue 50
COLOR_CALLOUT_TIP_BG = "F0FDF4"   # Green 50
COLOR_CALLOUT_WARN_BG = "FEF2F2"  # Red 50
COLOR_CALLOUT_NOTE_BG = "FFFBEB"  # Amber 50

RGB_PRIMARY = RGBColor(30, 58, 138)
RGB_SECONDARY = RGBColor(37, 99, 235)
RGB_DARK = RGBColor(15, 23, 42)
RGB_MUTED = RGBColor(71, 85, 105)
RGB_WHITE = RGBColor(255, 255, 255)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_left_border(cell, hex_color="2563EB", sz="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_table_borders(table, border_color=COLOR_BORDER_HEX):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def setup_document():
    doc = docx.Document()
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        s.different_first_page_header_footer = True
        
        # Setup Header for page 2+
        header = s.header
        hp = header.paragraphs[0]
        hp.text = "GlobalCo Job Portal — Master Technical Architecture & Interview Handbook"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Calibri"
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGB_MUTED
        
        # Setup Footer for page 2+
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.text = "Confidential | Master Technical Interview Preparation Notes | GlobalCo-JobBoard"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style.font.name = "Calibri"
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGB_MUTED

    # Base style
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGB_DARK
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    return doc

def add_cover_page(doc):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    
    # Title Box / Banner
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("GLOBALCO JOB PORTAL")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(32)
    r_title.font.bold = True
    r_title.font.color.rgb = RGB_PRIMARY
    p_title.paragraph_format.space_after = Pt(2)
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Master Project Architecture, Reverse-Engineered Deep Dive & Technical Interview Preparation Manual")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGB_SECONDARY
    r_sub.font.bold = True
    p_sub.paragraph_format.space_after = Pt(24)

    p_desc = doc.add_paragraph()
    r_desc = p_desc.add_run(
        "A rigorous, end-to-end technical reference reverse-engineering every component, REST API endpoint, "
        "database schema, Spring Security JWT filter, React state lifecycle, business algorithm, and "
        "interview defense strategy for the enterprise-grade GlobalCo recruitment platform."
    )
    r_desc.font.name = 'Calibri'
    r_desc.font.size = Pt(11)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGB_MUTED
    p_desc.paragraph_format.space_after = Pt(36)

    # Metadata Table
    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, COLOR_BORDER_HEX)
    
    meta_rows = [
        ("Project Name", "GlobalCo Job Portal & Recruitment Management System"),
        ("Primary Architecture", "Decoupled 3-Tier Layered Architecture (RESTful Spring Boot + React SPA)"),
        ("Backend Framework", "Spring Boot 3.4.3 | Java 17/21 | Spring Security 6 | Spring Data JPA"),
        ("Frontend Framework", "React 18.2 | Vite 7 | Tailwind CSS 3.4 | Axios | Lucide React"),
        ("Database & Persistence", "H2 In-Memory DB (with MySQL Connector Compatibility) | Hibernate ORM"),
        ("Security Mechanism", "Stateless JWT (HMAC-SHA256) | Role-Based Access Control (RBAC) | BCrypt"),
        ("Target Roles", "Full-Stack Software Engineer / Senior Backend Engineer / React Specialist"),
        ("Document Purpose", "Master Technical Notes for End-to-End System Defense & Interview Excellence")
    ]

    col_widths = [Inches(2.2), Inches(4.5)]
    for i, (k, v) in enumerate(meta_rows):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)
        set_cell_background(c0, COLOR_BG_LIGHT_HEX)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGB_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGB_DARK

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(40)
    r_conf = p_post.add_run("STRICTLY CONFIDENTIAL — PREPARED EXCLUSIVELY FOR TECHNICAL INTERVIEW PREPARATION")
    r_conf.font.size = Pt(8.5)
    r_conf.font.bold = True
    r_conf.font.color.rgb = RGB_MUTED
    p_post.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGB_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGB_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGB_DARK
    return p

def add_heading_4(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGB_MUTED
    return p

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGB_DARK
    r = p.add_run(text)
    r.font.color.rgb = RGB_DARK
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGB_DARK
    r = p.add_run(text)
    r.font.color.rgb = RGB_DARK
    return p

def add_callout(doc, text, title="NOTE", box_type="info"):
    bg_map = {
        "info": (COLOR_CALLOUT_INFO_BG, COLOR_SECONDARY_HEX),
        "tip": (COLOR_CALLOUT_TIP_BG, "10B981"),
        "warn": (COLOR_CALLOUT_WARN_BG, "EF4444"),
        "note": (COLOR_CALLOUT_NOTE_BG, "F59E0B")
    }
    bg_color, border_color = bg_map.get(box_type, (COLOR_CALLOUT_INFO_BG, COLOR_SECONDARY_HEX))

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.7)
    
    set_cell_background(cell, bg_color)
    set_cell_left_border(cell, hex_color=border_color, sz="36")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"[{title}] ")
    r_t.font.bold = True
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = RGB_PRIMARY

    r_c = p.add_run(text)
    r_c.font.size = Pt(9.5)
    r_c.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Listing: {caption}")
        r_cap.font.bold = True
        r_cap.font.size = Pt(9)
        r_cap.font.color.rgb = RGB_MUTED

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_background(cell, "F1F5F9")
    set_table_borders(table, "CBD5E1")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    lines = code_str.strip().split("\n")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, COLOR_BORDER_HEX)

    # Header Row
    hdr_row = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGB_WHITE

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg = COLOR_BG_LIGHT_HEX if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            r.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_qa(doc, question, short_ans, detailed_ans, project_conn, follow_up=None, follow_up_ans=None):
    add_heading_3(doc, f"Q: {question}")
    add_p(doc, short_ans, bold_prefix="30-Second Interview Response: ")
    add_p(doc, detailed_ans, bold_prefix="In-Depth Technical Explanation: ")
    add_p(doc, project_conn, bold_prefix="Exact Implementation in GlobalCo: ")
    if follow_up:
        add_p(doc, follow_up, bold_prefix="Interviewer Follow-Up: ")
        if follow_up_ans:
            add_p(doc, follow_up_ans, bold_prefix="Follow-Up Defense: ")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_screenshot(doc, img_path, caption):
    if not os.path.exists(img_path):
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(6.0))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(f"Figure: {caption}")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGB_MUTED
