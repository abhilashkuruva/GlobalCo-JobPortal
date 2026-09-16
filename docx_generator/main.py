import os
import sys

# Ensure root is in path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from docx_generator.styles import (
    setup_document, add_cover_page, add_heading_1, add_heading_2, 
    add_p, add_bullet, add_table, add_screenshot
)
from docx_generator.part1_overview_arch import generate_part_1
from docx_generator.part1b_foundations import generate_part_1b_foundations
from docx_generator.part2_backend_data import generate_part_2
from docx_generator.part2b_controllers_services import generate_part_2b_controllers_services
from docx_generator.part3_code_security import generate_part_3
from docx_generator.part3b_click_traces_failures import generate_part_3b_click_traces_failures
from docx_generator.part4_qa_defense import generate_part_4
from docx_generator.part4b_expanded_interview_bank import generate_part_4b_expanded_interview_bank
from docx_generator.part5_revision_cheatsheet import generate_part_5
from docx.shared import Inches, Pt

def build_complete_document():
    print("Initializing document styling and layout...")
    doc = setup_document()

    print("Generating Cover Page...")
    add_cover_page(doc)

    print("Generating Master Table of Contents...")
    add_heading_1(doc, "Master Table of Contents")
    add_p(doc, 
        "This master document is structured into 45+ comprehensive, exhaustive sections covering foundational computer science primers, "
        "system engineering, architecture, file-by-file code reverse-engineering, UI screen-by-screen visual documentation, "
        "button click traces, failure recovery, security, and technical interview defense:"
    )

    toc_sections = [
        ("Module 1", "System Overview & Foundations", "Executive Summary, Project Scope, Beginner Airport Analogy, Real-World Problem, Architectural Solution, User Roles & RBAC Matrix"),
        ("Module 2", "Core Architecture & Frontend", "Features Breakdown, Complete Directory Tree & Rationale, Technology Stack Deep Dive, 3-Tier System Architecture, Frontend Architecture"),
        ("Module 3", "Foundations Primer (Zero-to-Hero)", "Networking Mechanics, HTTP/HTTPS Anatomy, Java 21 & JVM Memory, Spring Boot IoC/DI, Bean Lifecycle, DispatcherServlet Pipeline, React Virtual DOM & Fiber, Relational Theory, Normalization (1NF-3NF), ACID Guarantees, Hibernate Internals, REST Principles, JWT Math, BCrypt Hashing"),
        ("Module 4", "Backend & API Specification", "Backend Architecture, Database Architecture (3NF), Entity Model & Data Dictionary, Complete 18+ REST API Reference Table"),
        ("Module 5", "Controllers & Services Deep Dive", "Exhaustive Inspection of All 14 Controllers, All 11 Domain Services, All 13 JPA Entities, All 13 Repositories, and DTO Validation Models"),
        ("Module 6", "Security & System Workflows", "Authentication Engineering (JWT), Authorization Filter Chain, Feature-by-Feature Deep Dive (6-Part Format), End-to-End Workflows"),
        ("Module 7", "Runtime & UI Screen Tour", "Runtime Analysis & Startup Lifecycle, Code-Level Deep Dive, 14 Screen-by-Screen UI Deep Dive with Embedded Screenshots, Algorithms & Business Logic"),
        ("Module 8", "Frontend Code & Failure Resiliency", "Frontend Component Hierarchy, AuthContext State, Axios Interceptor Pipeline, 20+ 'If I Click This Button' Execution Traces, 20+ 'What Happens If...' Edge Cases & Failure Scenarios"),
        ("Module 9", "Reliability, Testing & Cloud", "OWASP Top 10 Security, Error Handling Architecture, Configuration Deep Dive, Dependency Analysis, Automated Testing (11 Tests), Cloud Deployment"),
        ("Module 10", "Audits & Improvement Roadmap", "Audit Bugs Found & Fixed, Current Limitations, Phased 1M User Scalability Roadmap"),
        ("Module 11", "Master Interview Preparation", "30-Sec Pitch & 4-Tier Walkthroughs, Top 15 Technical Q&A, Difficult Follow-Ups, 'If Interviewer Opens My Code' Guide, How to Defend Project"),
        ("Module 12", "Expanded Technical Interview Bank", "Deep Dive Q&A on JVM, Constructor Injection, B-Tree Indexes, Optimistic vs Pessimistic Locking, React List Keys, useMemo vs useCallback, CORS, AuthN vs AuthZ, 1M User Scaling"),
        ("Module 13", "Learning Path & Revision", "Beginner-to-Master Learning Path (Levels 1-12), Technical Glossary (40+ Terms), Traceability Matrix, Cheat Sheet, One-Day Before Interview Revision Guide")
    ]
    
    toc_headers = ["Curriculum Module", "Theme / Focus Area", "Key Deliverables & Architectural Topics Included"]
    toc_rows = [[t[0], t[1], t[2]] for t in toc_sections]
    add_table(doc, toc_headers, toc_rows, [Inches(1.2), Inches(2.2), Inches(3.3)])
    doc.add_page_break()

    print("Building Part 1 (Sections 1 - 10: Overview, Architecture, Tech Stack)...")
    generate_part_1(doc)
    doc.add_page_break()

    print("Building Foundations Primer (Zero-to-Hero Computer Science Concepts)...")
    generate_part_1b_foundations(doc)

    print("Building Part 2 (Sections 11 - 18: Backend Architecture, APIs, Workflows)...")
    generate_part_2(doc)
    doc.add_page_break()

    print("Building Part 2b (Exhaustive Controllers, Services & Data Tier Deep Dive)...")
    generate_part_2b_controllers_services(doc)

    print("Building Part 3 (Sections 19 - 27: Runtime, Code Dive, 14 UI Screens)...")
    generate_part_3(doc)
    doc.add_page_break()

    print("Building Part 3b (Frontend State, 20+ Button Traces, 20+ Failure Scenarios)...")
    generate_part_3b_click_traces_failures(doc)

    print("Building Part 4 (Sections 28 - 34: Bugs, Roadmap, Interview Walkthroughs)...")
    generate_part_4(doc)
    doc.add_page_break()

    print("Building Part 4b (Expanded 5-Tier Master Technical Interview Bank)...")
    generate_part_4b_expanded_interview_bank(doc)

    print("Building Part 5 (Sections 35 - 39: Learning Path, Glossary, Revision Guide)...")
    generate_part_5(doc)

    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "Project_Master_Preparation_Notes.docx"))
    print(f"Saving compiled master document to: {output_path}")
    doc.save(output_path)
    print(f"SUCCESS: Document built successfully! File size: {os.path.getsize(output_path)} bytes")

    # Also mirror to project-named docx
    mirror_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "GlobalCo_JobBoard_Master_Preparation_Notes.docx"))
    doc.save(mirror_path)
    print(f"SUCCESS: Mirrored document saved to: {mirror_path}")

    return output_path

if __name__ == "__main__":
    build_complete_document()
